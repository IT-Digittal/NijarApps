"""Anexos I y II de suministros eléctricos (formato de los pliegos): lógica pura."""

from __future__ import annotations

import io
from datetime import date, datetime
from types import SimpleNamespace

import openpyxl

from nijar_dti.connectors import datadis as dd
from nijar_dti.services import energia_anexos_service as ax

MADRID = ax.MADRID


def _h(y, m, d, h):
    return datetime(y, m, d, h, tzinfo=MADRID)


class TestPeriodos:
    def test_20td(self):
        assert ax.periodo_tarifario(_h(2026, 10, 27, 11), "2.0TD") == "P1"  # martes 11 h
        assert ax.periodo_tarifario(_h(2026, 10, 27, 9), "2.0TD") == "P2"
        assert ax.periodo_tarifario(_h(2026, 10, 27, 23), "2.0TD") == "P2"
        assert ax.periodo_tarifario(_h(2026, 10, 27, 3), "2.0TD") == "P3"
        assert ax.periodo_tarifario(_h(2026, 10, 25, 11), "2.0TD") == "P3"  # domingo

    def test_30td_por_temporada(self):
        assert ax.periodo_tarifario(_h(2026, 1, 13, 10), "3.0TD") == "P1"  # enero, alta
        assert ax.periodo_tarifario(_h(2026, 1, 13, 15), "3.0TD") == "P2"
        assert ax.periodo_tarifario(_h(2026, 3, 10, 10), "3.0TD") == "P2"  # marzo, media-alta
        assert ax.periodo_tarifario(_h(2026, 6, 9, 10), "3.0TD") == "P3"  # junio, media
        assert ax.periodo_tarifario(_h(2026, 4, 14, 10), "6.1TD") == "P4"  # abril, baja
        assert ax.periodo_tarifario(_h(2026, 4, 14, 8), "6.1TD") == "P5"
        assert ax.periodo_tarifario(_h(2026, 4, 14, 7), "3.0TD") == "P6"

    def test_festivo_nacional_es_valle(self):
        assert ax.periodo_tarifario(_h(2026, 12, 8, 11), "3.0TD") == "P6"  # martes festivo
        assert ax.periodo_tarifario(_h(2026, 12, 8, 11), "2.0TD") == "P3"


class TestDatos:
    def test_detalle_contrato_toma_el_ultimo(self):
        datos = {
            "contract": [
                {"startDate": "2020/01/01", "accessFare": "2.0A", "contractedPowerkW": [5]},
                {
                    "startDate": "2023/06/01",
                    "accessFare": "3.0TD",
                    "tension": "400",
                    "contractedPowerkW": [15, 15, 15, 15, 15, 20],
                    "distributor": "E-DISTRIBUCION REDES DIGITALES",
                },
            ]
        }
        c = dd.detalle_contrato(datos)
        assert c["tarifa_acceso"] == "3.0TD" and c["tension"] == "400"
        assert c["potencias_kw"] == [15, 15, 15, 15, 15, 20]

    def test_fila_anexo_i(self):
        s = SimpleNamespace(
            cups="ES0031000000000001AB",
            edificio="CL REAL 1",
            tipo="alumbrado",
            metadatos={
                "fuente": "datadis",
                "direccion": "CL REAL 1, NIJAR",
                "provincia": "Almería",
                "municipio": "Níjar",
                "esencial": True,
                "contrato": {
                    "tarifa_acceso": "2.0TD",
                    "potencias_kw": [3.45, 3.45],
                    "distribuidora": "E-DISTRIBUCION REDES DIGITALES",
                },
            },
        )
        f = ax.fila_anexo_i(s)  # type: ignore[arg-type]
        assert f.peaje == "2.0TD" and f.esencial == "SI"
        assert f.potencias == [3.45, 3.45, None, None, None, None]
        assert f.cif_distribuidora == "B82846817"

    def test_distribuidora_desconocida_queda_en_blanco(self):
        assert ax.cif_distribuidora("ELECTRICA LOCAL SL") == ("", "ELECTRICA LOCAL SL")

    def test_ultimos_12_meses(self):
        assert ax.ultimos_12_meses(date(2026, 10, 27)) == (date(2025, 10, 1), date(2026, 9, 30))
        assert ax.ultimos_12_meses(date(2026, 1, 5)) == (date(2025, 1, 1), date(2025, 12, 31))


def _filas():
    f = ax.FilaAnexoI(
        "ES1",
        "Dir",
        "Níjar",
        "Almería",
        "2.0TD",
        "230",
        "",
        [3.45] + [None] * 5,
        "B82846817",
        "E-Distribución",
        "alumbrado",
    )
    g = ax.FilaAnexoII(
        "2.0TD", 1, {"P1": 10.5, "P2": 5.0, "P3": 20.0, "P4": 0.0, "P5": 0.0, "P6": 0.0}
    )
    return [f], [g], {"ES1": g.energia}


class TestFicheros:
    def test_csv_con_formato_espanol(self):
        fi, fii, _ = _filas()
        txt = ax.csv_anexo("ii", fi, fii)
        assert txt.startswith("﻿") and "T1 (kWh)" in txt
        assert "1;2.0TD;10,50;5,00;20,00" in txt
        assert ax.csv_anexo("i", fi, fii).splitlines()[1].startswith("ES1;Dir;Níjar")

    def test_excel_tiene_los_dos_anexos(self):
        fi, fii, pc = _filas()
        libro = openpyxl.load_workbook(
            io.BytesIO(ax.excel_anexos(fi, fii, pc, date(2025, 10, 1), date(2026, 9, 30)))
        )
        assert libro.sheetnames[:2] == ["Anexo I - CUPS", "Anexo II - Consumos"]
        h2 = libro["Anexo II - Consumos"]
        filas = [r for r in h2.iter_rows(values_only=True) if r and r[1] in ("2.0TD", "TOTAL")]
        assert filas[0][:3] == (1, "2.0TD", 10.5) and filas[-1][1] == "TOTAL"
