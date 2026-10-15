"""Lectura del fichero que descarga datadis.es → Datos públicos → Consultas.

Mientras no hay usuario de Datadis para su API, la tabla «Consumos y
contratos» se puede descargar a mano desde la web (botón «Descargar», un mes
por consulta) y subirla a la plataforma. Este módulo la convierte en los mismos
``RegistroPublico`` que produce la API.

Formato esperado (columnas de la web): Comunidad · Provincia · Municipio ·
Fecha · Sector · Tensión · Contratos (Suma) · Energía (Suma kWh). Se aceptan CSV
(«;», «,» o tabulador; UTF-8 o Latin-1) y Excel (.xlsx), con o sin líneas de
texto antes de la cabecera, cabeceras en español o inglés y números con formato
español («191.589», «5.311,151»). Si la energía viene en MWh, se pasa a kWh.
"""

from __future__ import annotations

import csv
import io
from datetime import date, datetime
from typing import Any

from nijar_dti.connectors.datadis import RegistroPublico, normalizar, parsear_registro_publico

# Columna del fichero (normalizada, empieza por…) → clave que entiende el parser de la API
_COLUMNAS: list[tuple[tuple[str, ...], str]] = [
    (("COMUNIDAD", "COMMUNITY", "AUTONOMOUS"), "community"),
    (("PROVINCIA", "PROVINCE"), "province"),
    (("MUNICIPIO", "MUNICIPALITY", "LOCALIDAD"), "municipality"),
    (("FECHA", "DATE", "DIA", "DAY"), "date"),
    (("SECTOR", "ECONOMIC"), "economicSector"),
    (("TENSION", "TARIFA", "VOLTAGE", "FARE"), "tension"),
    (("CONTRATO", "CONTRACT", "N CONTRATOS"), "sumContracts"),
    (("ENERGIA", "ENERGY", "CONSUMO", "CONSUMPTION"), "sumEnergy"),
]


class FicheroDatadisError(ValueError):
    """El fichero no tiene el formato de la consulta de Datadis."""


def _clave(cabecera: str) -> str | None:
    c = normalizar(cabecera).replace("º", "").replace("Nº", "N").strip()
    for prefijos, clave in _COLUMNAS:
        if any(c.startswith(p) for p in prefijos):
            return clave
    return None


def _celda(v: Any) -> Any:
    if isinstance(v, datetime):
        return v.strftime("%Y/%m/%d")
    if isinstance(v, date):
        return v.strftime("%Y/%m/%d")
    return v


def _filas_csv(contenido: bytes) -> list[list[Any]]:
    for codif in ("utf-8-sig", "latin-1"):
        try:
            texto = contenido.decode(codif)
            break
        except UnicodeDecodeError:
            continue
    muestra = "\n".join(texto.splitlines()[:30])
    try:
        delimitador = csv.Sniffer().sniff(muestra, delimiters=";,\t").delimiter
    except csv.Error:
        delimitador = ";" if muestra.count(";") >= muestra.count(",") else ","
    return list(csv.reader(io.StringIO(texto), delimiter=delimitador))


def _filas_xlsx(contenido: bytes) -> list[list[Any]]:
    from openpyxl import load_workbook  # type: ignore[import-untyped]

    libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    hoja = libro.worksheets[0]
    return [[_celda(c) for c in fila] for fila in hoja.iter_rows(values_only=True)]


def leer_fichero(contenido: bytes, nombre: str = "") -> list[RegistroPublico | None]:
    """Convierte el fichero descargado de Datadis en registros (``None`` = fila ilegible)."""
    es_xlsx = contenido[:2] == b"PK" or nombre.lower().endswith((".xlsx", ".xlsm"))
    filas = _filas_xlsx(contenido) if es_xlsx else _filas_csv(contenido)

    # La cabecera es la primera fila que tenga, al menos, fecha y energía
    cab_idx, mapa = None, {}
    for i, fila in enumerate(filas[:50]):
        claves = {j: _clave(str(v)) for j, v in enumerate(fila) if v not in (None, "")}
        if {"date", "sumEnergy"} <= set(claves.values()):
            cab_idx, mapa = i, {j: k for j, k in claves.items() if k}
            break
    if cab_idx is None:
        raise FicheroDatadisError(
            "No se encuentra la cabecera (Fecha, Sector, Tensión, Contratos, Energía) "
            "de la consulta «Consumos y contratos» de datadis.es"
        )
    en_mwh = any("MWH" in normalizar(str(v)) for v in filas[cab_idx] if v)

    registros: list[RegistroPublico | None] = []
    for fila in filas[cab_idx + 1 :]:
        if not any(v not in (None, "") for v in fila):
            continue
        datos = {mapa[j]: fila[j] for j in mapa if j < len(fila)}
        reg = parsear_registro_publico(datos)
        if reg is not None and en_mwh:
            reg = RegistroPublico(**{**reg.__dict__, "energia_kwh": reg.energia_kwh * 1000})
        registros.append(reg)
    return registros
